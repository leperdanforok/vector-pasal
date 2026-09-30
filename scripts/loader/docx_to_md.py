"""Convert .docx files in data/raw_pdf/ to Markdown transcriptions in data/transcriptions/.

Usage:
    python scripts/loader/docx_to_md.py

Produces one .md file per .docx, named by the same slug convention that
load_perda_bolmong.py uses (lowercase, spaces/underscores → hyphens).
Skips files that already have a transcription.
"""

from pathlib import Path

try:
    import docx
except ImportError:
    print("python-docx not installed. Run:")
    print("  .\\venv\\Scripts\\pip.exe install python-docx")
    raise SystemExit(1)

BASE_DIR = Path(__file__).parent.parent.parent
RAW_DIR = BASE_DIR / "data" / "raw_pdf"
OUT_DIR = BASE_DIR / "data" / "transcriptions"
OUT_DIR.mkdir(parents=True, exist_ok=True)


def docx_to_markdown(path: Path) -> str:
    doc = docx.Document(path)
    lines: list[str] = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            lines.append("")
            continue

        style = (para.style.name if para.style else "") or ""
        style = style.lower()
        if "heading 1" in style:
            lines.append(f"# {text}")
        elif "heading 2" in style:
            lines.append(f"## {text}")
        elif "heading 3" in style:
            lines.append(f"### {text}")
        elif "heading 4" in style:
            lines.append(f"#### {text}")
        else:
            lines.append(text)

    for table in doc.tables:
        lines.append("")
        for i, row in enumerate(table.rows):
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            lines.append("| " + " | ".join(cells) + " |")
            if i == 0:
                lines.append("| " + " | ".join(["---"] * len(cells)) + " |")
        lines.append("")

    return "\n".join(lines)


def main():
    docx_files = sorted(RAW_DIR.glob("*.docx"))
    print(f"Found {len(docx_files)} .docx files in {RAW_DIR}\n")

    if not docx_files:
        print("No .docx files found.")
        return

    for path in docx_files:
        slug = path.stem.lower().replace(" ", "-").replace("_", "-")
        out_path = OUT_DIR / f"{slug}.md"

        if out_path.exists():
            print(f"  SKIP {path.name} → transcription already exists")
            continue

        print(f"  Converting: {path.name}")
        md = docx_to_markdown(path)
        out_path.write_text(md, encoding="utf-8")
        print(f"    → {out_path.name} ({len(md)} chars)")

    print("\nDone. Run load_perda_bolmong.py to ingest.")


if __name__ == "__main__":
    main()
